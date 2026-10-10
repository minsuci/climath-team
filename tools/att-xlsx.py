# 예비고1 출결현황 엑셀 — nas-daily-digest.mjs 가 JSON 을 넘기면 xlsx 로 쓴다 (2026-10-10, 원장님 #59).
#   python tools/att-xlsx.py <in.json> <out.xlsx>
# 받는 사람은 박솔 대리(수강료 · ERP 출결). 값만 쓴다(수식 없음) — 그 주의 사본이라 다시 계산할 것이 없고,
# 수식으로 쓰면 엑셀로 열기 전에는 값이 비어 보여 AI·미리보기가 «0» 으로 읽는다.
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
DOW = "월화수목금토일"
FILL = {
    "결": "F8D7DA", "지": "FFE5CC", "조": "FFE5CC",
    "?": "FFF3B0", "미보고": "FFF3B0", "밖": "E4E4E4",
}
HEAD = PatternFill("solid", fgColor="1F3A5F")
THIN = Side(style="thin", color="C8C8C8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def md(d, dow=True):
    y, m, dd = map(int, d.split("-"))
    s = f"{m}/{dd}"
    return s + f"({DOW[date(y, m, dd).weekday()]})" if dow else s


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


def head_row(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def body(c, center=False, mark=None):
    c.font = font()
    c.border = BOX
    c.alignment = Alignment(horizontal="center" if center else "left", vertical="center")
    if mark in FILL:
        c.fill = PatternFill("solid", fgColor=FILL[mark])


def sheet_name(name, used):
    s = "".join("_" if ch in '[]:*?/\\' else ch for ch in name)[:28]
    base, n = s, 2
    while s in used:
        s = f"{base[:26]}{n}"
        n += 1
    used.add(s)
    return s


def main(src, out):
    o = json.load(open(src, encoding="utf-8"))
    month = int(o["month"][5:])
    wb = Workbook()

    # ---- 읽는 법 ----
    ws = wb.active
    ws.title = "읽는 법"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 90
    title = f"예비고1 출결현황 — {month}월" + (" 확정본" if o.get("final") else " 누적")
    ws["A1"] = title
    ws["A1"].font = font(size=15, bold=True)
    lines = [
        ("기간", f"{md(o['from'])} ~ {md(o['until'])}" + ("  (그 달 끝까지)" if o.get("final") else "  (만든 날 전날까지)")
                + ("  · 업무보고는 9/14부터라 그 전 날은 없다" if o["from"][8:] != "01" else "")),
        ("만든 날", md(o["made"]) + " · 팀체크 업무보고에서 자동으로 옮김 (한민수)"),
        ("기준", "선생님이 업무보고에 찍은 출결. 수업관리 앱 체크인 기록이 아니다."),
        ("대상", "예비고1 정규반 전원 + 개진반의 중3 학생. 개진반은 반이 따로라 같은 학생이 두 줄일 수 있다(수강료가 따로다)."),
        ("", ""),
        ("출 · 결 · 지 · 조", "출석 · 결석 · 지각 · 조퇴 — 업무보고 그대로"),
        ("?", "그 날 보고는 냈는데 이 학생 출결을 안 찍었다 → 담임 선생님께 확인"),
        ("미보고", "수업 날인데 그 반 업무보고가 없다 → 담임 선생님께 확인 (출석으로 읽지 말 것)"),
        ("밖", "업무보고를 안 받는 반·요일이다(원장님 개진반 · 중등관에 보고하는 요일). 이 칸은 따로 챙긴다"),
        ("빈칸", "그 학생이 그 날 수업이 아니다(개진반 요일 · 등록 전 · 퇴원 뒤)"),
        ("", ""),
        ("요약 «수업»", "그 달 수업이 있던 날 수(«밖» 은 빼고). «출석» 에는 지각·조퇴도 든다."),
        ("요약 «확인 필요»", "? + 미보고. 0이 아니면 수강료를 정하기 전에 확인한다."),
        ("명단 변동", "업무보고에 적힌 학생 추가 · 퇴원 · 반 이동. 수강료 시작·끝 날짜를 여기서 본다."),
    ]
    outside = [c["name"] + " (" + c["teachers"] + ")" for c in o["classes"] if c.get("outside")]
    if outside:
        lines.append(("업무보고 밖 반", " · ".join(outside)))
    for i, (k, v) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=k).font = font(bold=True)
        ws.cell(row=i, column=2, value=v).font = font()
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if k in FILL:
            ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=FILL[k])

    # ---- 요약 ----
    sm = wb.create_sheet("요약")
    cols = ["반", "담임", "학생", "학교", "수업", "출석", "결석", "지각·조퇴", "확인 필요", "업무보고 밖", "보강 · 메모"]
    head_row(sm, 1, cols, [18, 10, 10, 14, 7, 7, 7, 9, 9, 10, 40])
    r = 2
    for c in o["classes"]:
        for row in c["rows"]:
            s = row["sum"]
            vals = [c["name"], c["teachers"], row["name"], row.get("school", ""), s["held"], s["att"], s["abs"], s["late"],
                    s["unknown"], s["out"], " · ".join(row.get("notes") or [])]
            for i, v in enumerate(vals, 1):
                cell = sm.cell(row=r, column=i, value=v)
                body(cell, center=5 <= i <= 10)
            if s["abs"]:
                sm.cell(row=r, column=7).fill = PatternFill("solid", fgColor=FILL["결"])
            if s["unknown"]:
                sm.cell(row=r, column=9).fill = PatternFill("solid", fgColor=FILL["?"])
            if s["out"]:
                sm.cell(row=r, column=10).fill = PatternFill("solid", fgColor=FILL["밖"])
            r += 1
    sm.freeze_panes = "C2"
    sm.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(r - 1, 1)}"

    # ---- 반마다 ----
    used = {"읽는 법", "요약", "명단 변동"}
    for c in o["classes"]:
        ws = wb.create_sheet(sheet_name(c["name"], used))
        dates = c["dates"]
        hdr = ["학생", "학교"] + [md(d) for d in dates] + ["수업", "출석", "결석", "확인"]
        head_row(ws, 1, hdr, [10, 14] + [8.5] * len(dates) + [6, 6, 6, 6])
        for i, row in enumerate(c["rows"], 2):
            body(ws.cell(row=i, column=1, value=row["name"]))
            body(ws.cell(row=i, column=2, value=row.get("school", "")))
            for j, d in enumerate(dates, 3):
                m = row["cells"].get(d, "")
                body(ws.cell(row=i, column=j, value=m), center=True, mark=m)
            s, k = row["sum"], 3 + len(dates)
            for j, v in enumerate([s["held"], s["att"], s["abs"], s["unknown"]]):
                body(ws.cell(row=i, column=k + j, value=v), center=True)
        ws.freeze_panes = "C2"
        ws.row_dimensions[1].height = 30
        note = f"{c['teachers']} · 수업 {len(dates)}일 · {len(c['rows'])}명" + (" · 업무보고 밖 — 이 반은 따로 챙긴다" if c.get("outside") else "")
        ws.cell(row=len(c["rows"]) + 3, column=1, value=note).font = font(color="666666")

    # ---- 명단 변동 ----
    mv = wb.create_sheet("명단 변동")
    head_row(mv, 1, ["보고한 날", "선생님", "종류", "학생", "내용", "언제부터", "메모"], [11, 10, 10, 10, 34, 11, 40])
    if not o["moves"]:
        mv.cell(row=2, column=1, value="이 기간 업무보고에 적힌 명단 변동이 없다.").font = font(color="666666")
    for i, m in enumerate(o["moves"], 2):
        vals = [md(m["reported"]), m["teacher"], m["kind"], m["name"], m["what"], md(m["date"]) if m.get("date") else "", m.get("note", "")]
        for j, v in enumerate(vals, 1):
            body(mv.cell(row=i, column=j, value=v), center=j in (1, 3, 6))

    wb.move_sheet("명단 변동", offset=2 - (len(wb.sheetnames) - 1))   # 읽는 법 · 요약 · 명단 변동 · 반들
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
